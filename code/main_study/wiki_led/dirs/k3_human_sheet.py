#!/usr/bin/env python3
"""Addendum k, K3: build the 100-page human check sheet of the Test F AI sorting (prereg/addenda_wiki_led/2026-10-22_k_intent_dates_human.json).
Sample (seed 20261022): 80 pages at random among universe pages with at least one theme, 20 at random among universe pages sorted to none;
left out: the judge's 400 pages, the 250 pilot pages (lines 2 to 251 of pages_chunk_05.tsv) and the three instruction examples.
The rater sees title, short description, first sentence and the AI's themes (name and broader group); not blind (record).
Choices not fixed by the record: the sheet shows the themes in the AI's order without the confidence numbers (so the rater is not led by
them; confidence is kept in the key); the page order in the sheet is shuffled with the same seed.
Writes data/db/wiki_led/dirs/k3_human/k3_sample_key.csv and paper/K3_Human_Check_Sheet_<date>.xlsx.
Run from code/:  ../.venv/bin/python main_study/wiki_led/dirs/k3_human_sheet.py"""
import csv, glob, pathlib, time
import numpy as np, pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.datavalidation import DataValidation

ROOT = pathlib.Path(__file__).resolve().parents[4]
T = ROOT / 'data' / 'db' / 'wiki_led' / 'testF'
OUT = ROOT / 'data' / 'db' / 'wiki_led' / 'dirs' / 'k3_human'
SEED = 20261022
guide = {r['code']: r for r in csv.DictReader(open(ROOT / 'prereg' / 'addenda_wiki_led' / '2026-10-21_g_theme_guide.tsv'), delimiter='\t')}
items = pd.read_parquet(T / 'items_pages.parquet')
A = pd.read_parquet(T / 'assign_pages.parquet').sort_values(['id', 'rank'])
text = pd.concat([pd.read_csv(f, sep='\t', dtype=str, quoting=csv.QUOTE_NONE, keep_default_na=False) for f in sorted(glob.glob(str(T / 'pages_chunk_*.tsv')))])
text = text.set_index('id')
judge = set(pd.read_csv(T / 'judge_pages.tsv', sep='\t', dtype=str, quoting=csv.QUOTE_NONE, keep_default_na=False).id)
pilot = set(pd.read_csv(T / 'pages_chunk_05.tsv', sep='\t', dtype=str, quoting=csv.QUOTE_NONE, keep_default_na=False).id.iloc[:250])
U = items[(items.set == 'universe') & ~items.prompt_example.astype(bool) & ~items.id.isin(judge | pilot)]
rng = np.random.default_rng(SEED)
themed = U[U.status == 'themes'].id.to_numpy(); none = U[U.status == 'none'].id.to_numpy()
pick = list(rng.choice(np.sort(themed), 80, replace=False)) + list(rng.choice(np.sort(none), 20, replace=False))
order = list(rng.permutation(len(pick)))
rows = []
for n, k in enumerate(order, 1):
    pid = pick[k]
    t = text.loc[pid]
    a = A[A.id == pid]
    th = [(r.code, int(r.conf)) for r in a.itertuples()]
    rows.append(dict(n=n, id=pid, title=t.title, description=t.description, first_sentence=t.first_sentence,
                     status='themes' if th else 'none', **{f'code{i+1}': c for i, (c, _) in enumerate(th)},
                     **{f'conf{i+1}': cf for i, (_, cf) in enumerate(th)}))
key = pd.DataFrame(rows)
key.to_csv(OUT / 'k3_sample_key.csv', index=False)

wb = Workbook()
ws0 = wb.active; ws0.title = 'How to do it'
howto = [
    'Human check of the AI sorting (about an hour). Khant Phyo Wai, student ID 01323050.',
    '',
    'Each row is one Wikipedia page from Test F. Read its title, short description and first sentence.',
    'The AI was asked: which themes of library books would a reader interested in this page naturally look for?',
    'For each theme the AI gave, choose in the drop-down:  fits  /  partly  /  does not fit.',
    'Then say whether a clearly better theme is missing:  yes / no.',
    'Rows where the AI gave no theme: say whether you agree that no theme fits:  agree / a theme fits.',
    'Use your own knowledge of the topic; there is no need to look anything up. Notes are optional.',
    'Save the file when done and tell Claude; nothing else is needed.',
]
for i, line in enumerate(howto, 1):
    ws0.cell(i, 1, line).font = Font(bold=(i == 1), size=12)
ws0.column_dimensions['A'].width = 120
ws = wb.create_sheet('Pages')
head = ['#', 'Page title', 'Short description', 'First sentence', 'AI theme 1', 'Theme 1 rating', 'AI theme 2', 'Theme 2 rating',
        'AI theme 3', 'Theme 3 rating', 'Better theme missing?', 'No-theme pages: agree?', 'Note']
ws.append(head)
for c in ws[1]:
    c.font = Font(bold=True); c.fill = PatternFill('solid', fgColor='DDE7F0'); c.alignment = Alignment(wrap_text=True, vertical='top')
rate = DataValidation(type='list', formula1='"fits,partly,does not fit"', allow_blank=True)
yn = DataValidation(type='list', formula1='"yes,no"', allow_blank=True)
agree = DataValidation(type='list', formula1='"agree,a theme fits"', allow_blank=True)
for dv in (rate, yn, agree):
    ws.add_data_validation(dv)
grey = PatternFill('solid', fgColor='EEEEEE')
for r in rows:
    def name(i):
        c = r.get(f'code{i}')
        return f"{guide[c]['theme']}  ({guide[c]['broader group']})" if isinstance(c, str) and c in guide else ''
    ws.append([r['n'], r['title'], r['description'], r['first_sentence'], name(1), '', name(2), '', name(3), '', '', '', ''])
    i = ws.max_row
    for j, col in ((1, 'F'), (2, 'H'), (3, 'J')):
        if name(j):
            rate.add(f'{col}{i}')
        else:
            ws[f'{col}{i}'].fill = grey
    if r['status'] == 'themes':
        yn.add(f'K{i}'); ws[f'L{i}'].fill = grey
    else:
        agree.add(f'L{i}'); ws[f'K{i}'].fill = grey
        ws[f'E{i}'] = '(the AI gave no theme)'
    for c in ws[i]:
        c.alignment = Alignment(wrap_text=True, vertical='top')
for col, w in zip('ABCDEFGHIJKLM', (5, 28, 30, 60, 30, 14, 30, 14, 30, 14, 14, 18, 30)):
    ws.column_dimensions[col].width = w
ws.freeze_panes = 'B2'
f = ROOT / 'paper' / f'K3_Human_Check_Sheet_{time.strftime("%Y-%m-%d_%H%M")}.xlsx'
wb.save(f)
print('pages', len(rows), '| themed', int((key.status == "themes").sum()), '| none', int((key.status == "none").sum()), '| themes to rate',
      int(key[[c for c in key.columns if c.startswith('code')]].notna().sum().sum()), '|', f)
